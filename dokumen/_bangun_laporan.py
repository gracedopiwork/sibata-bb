# -*- coding: utf-8 -*-
"""Susun Laporan Laboratorium Prakom SIBATA-BB (Bab I–IV) ke Word."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.shared import Cm, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent
GAMBAR = ROOT / "gambar" / "bukti"
OUT = ROOT / "Laporan Laboratorium Prakom SIBATA-BB (Lengkap).docx"

NAVY = RGBColor(11, 31, 58)
GOLD = RGBColor(201, 162, 39)
INK = RGBColor(22, 22, 22)
WHITE = RGBColor(255, 255, 255)


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), kwargs.get("val", "single"))
        el.set(qn("w:sz"), kwargs.get("sz", "8"))
        el.set(qn("w:color"), kwargs.get("color", "1B1F2A"))
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_run(run, size=12, bold=False, italic=False, color=INK, font="Times New Roman"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def para_fmt(p, align="justify", first=True, before=0, after=6, space=1.5, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = space
    if keep:
        pf.keep_with_next = True
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.first_line_indent = Cm(1.25) if first else Cm(0)


def add_text(doc, text, *, size=12, bold=False, italic=False, align="justify", first=True, before=0, after=6, space=1.5, color=INK):
    p = doc.add_paragraph()
    para_fmt(p, align=align, first=first, before=before, after=after, space=space)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_bab(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, align="center", first=False, before=18, after=12, space=1.5)
    r = p.add_run(text)
    set_run(r, size=14, bold=True, color=NAVY)
    return p


def add_huruf(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, align="left", first=False, before=12, after=6, space=1.5, keep=True)
    r = p.add_run(text)
    set_run(r, size=12, bold=True, color=NAVY)
    return p


def add_angka(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, align="left", first=False, before=8, after=4, space=1.5, keep=True)
    r = p.add_run(text)
    set_run(r, size=12, bold=True, color=NAVY)
    return p


def add_item(doc, text, hang=True):
    p = doc.add_paragraph()
    para_fmt(p, align="justify", first=False, before=0, after=4, space=1.5)
    if hang:
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(-0.5)
    r = p.add_run(text)
    set_run(r, size=12)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, align="center", first=False, before=4, after=12, space=1.0)
    r = p.add_run(text)
    set_run(r, size=11, italic=True, color=NAVY)
    return p


def table_caption(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, align="center", first=False, before=10, after=4, space=1.0, keep=True)
    r = p.add_run(text)
    set_run(r, size=11, bold=True, color=NAVY)
    return p


def add_figure(doc, filename: str, cap: str, width=15.8):
    path = GAMBAR / filename
    p = doc.add_paragraph()
    para_fmt(p, align="center", first=False, before=8, after=2, space=1.0)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width))
    caption(doc, cap)


def set_table_full(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    ic = tblPr.find(qn("w:tblW"))
    if ic is None:
        ic = OxmlElement("w:tblW")
        tblPr.append(ic)
    ic.set(qn("w:w"), "5000")
    ic.set(qn("w:type"), "pct")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "1B1F2A")
        borders.append(el)
    tblPr.append(borders)


def fill_table(table, rows, header=True, sizes=None, center_last=False):
    set_table_full(table)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or (center_last and j == len(row) - 1)) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(val))
            is_head = header and i == 0
            set_run(r, size=10, bold=is_head, color=WHITE if is_head else INK)
            if is_head:
                shade(cell, "0B1F3A")
            elif i % 2 == 0:
                shade(cell, "F3F5F8")
            set_cell_border(cell)
    if sizes:
        for row in table.rows:
            for j, w in enumerate(sizes):
                row.cells[j].width = Cm(w)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run(run, size=10, color=NAVY)


def setup_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.5)
    sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = INK
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    sec.different_first_page_header_footer = True
    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("Laporan Laboratorium Prakom  ·  SIBATA-BB  ·  Kejari Wajo")
    set_run(hr, size=9, italic=True, color=NAVY)
    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Pelatihan Fungsional Penguatan Pranata Komputer Kategori Keahlian  ·  2026  ·  halaman ")
    set_run(fr, size=9, color=NAVY)
    add_page_number(fp)
    return doc


def cover(doc):
    add_text(doc, "KEJAKSAAN AGUNG REPUBLIK INDONESIA", size=13, bold=True, align="center", first=False, before=24, after=0, space=1.15, color=NAVY)
    add_text(doc, "KEJAKSAAN TINGGI SULAWESI SELATAN", size=12, bold=True, align="center", first=False, before=0, after=0, space=1.15, color=NAVY)
    add_text(doc, "KEJAKSAAN NEGERI WAJO", size=12, bold=True, align="center", first=False, before=0, after=6, space=1.15, color=NAVY)
    p = doc.add_paragraph()
    para_fmt(p, align="center", first=False, before=0, after=12)
    r = p.add_run("━" * 42)
    set_run(r, size=12, color=GOLD)

    add_text(doc, "LAPORAN LABORATORIUM PRANATA KOMPUTER", size=16, bold=True, align="center", first=False, before=18, after=4, space=1.15, color=NAVY)
    add_text(doc, "PELATIHAN FUNGSIONAL PENGUATAN PRANATA KOMPUTER", size=12, bold=True, align="center", first=False, before=0, after=0, space=1.15)
    add_text(doc, "KATEGORI KEAHLIAN", size=12, bold=True, align="center", first=False, before=0, after=0, space=1.15)
    add_text(doc, "TAHUN ANGGARAN 2026", size=12, bold=True, align="center", first=False, before=0, after=18, space=1.15)

    add_text(doc, "JUDUL INOVASI", size=11, bold=True, align="center", first=False, before=12, after=6, color=NAVY)
    add_text(doc, "Sistem Informasi Barang Bukti dan Barang Rampasan Terintegrasi Bot Telegram dan Portal Web Labeling QR Code (SIBATA-BB)", size=14, bold=True, align="center", first=False, before=0, after=6, space=1.3, color=NAVY)
    add_text(doc, "pada Seksi Pengelolaan Barang Bukti dan Barang Rampasan (PB3R)\nKejaksaan Negeri Wajo", size=12, align="center", first=False, before=0, after=24, space=1.3)

    add_text(doc, "Oleh", size=12, align="center", first=False, before=24, after=4)
    add_text(doc, "Grace Yoby Dopi", size=14, bold=True, align="center", first=False, before=0, after=0, color=NAVY)
    add_text(doc, "NIP 199710152022032004", size=12, align="center", first=False, before=0, after=36)

    add_text(doc, "KEJAKSAAN AGUNG", size=13, bold=True, align="center", first=False, before=48, after=0, color=NAVY)
    add_text(doc, "TAHUN 2026", size=12, bold=True, align="center", first=False, before=0, after=0, color=NAVY)
    doc.add_page_break()


def toc(doc):
    add_bab(doc, "DAFTAR ISI")
    items = [
        ("HALAMAN JUDUL", "1"),
        ("DAFTAR ISI", "2"),
        ("DAFTAR TABEL", "3"),
        ("DAFTAR GAMBAR", "4"),
        ("BAB I  PENDAHULUAN", "5"),
        ("     A. Latar Belakang", "5"),
        ("     B. Tujuan", "7"),
        ("     C. Manfaat", "7"),
        ("BAB II  TELAAH MASALAH DAN RENCANA KEGIATAN TI", "8"),
        ("     A. Identifikasi dan Analisis Masalah/Isu TI", "9"),
        ("     B. Deskripsi Kegiatan dan Solusi", "10"),
        ("     C. Pemetaan Kompetensi Prakom", "10"),
        ("     D. Penjadwalan (Gantt Chart)", "11"),
        ("BAB III  HASIL KEGIATAN DAN BUKTI KEGIATAN", "12"),
        ("     A. Pelaksanaan dan Hasil Kegiatan Prakom", "12"),
        ("          1. Perancangan Sistem Informasi SIBATA-BB", "12"),
        ("          2. Pembuatan Sistem Informasi SIBATA-BB", "18"),
        ("          3. Testing / Pengujian Sistem Informasi SIBATA-BB", "23"),
        ("     B. Dampak Hasil Kegiatan", "26"),
        ("BAB IV  KENDALA DAN RENCANA TINDAK LANJUT", "27"),
        ("     A. Kendala", "27"),
        ("     B. Rencana Tindak Lanjut", "28"),
    ]
    for title, page in items:
        p = doc.add_paragraph()
        para_fmt(p, align="left", first=False, before=0, after=2, space=1.3)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(f"{title}\t{page}")
        bold = title.startswith("BAB") or title.startswith("DAFTAR") or title.startswith("HALAMAN")
        set_run(r, size=12, bold=bold, color=NAVY if bold else INK)
    doc.add_page_break()

    add_bab(doc, "DAFTAR TABEL")
    tables = [
        "Tabel 2.1  Gap analisis isu TI SIBATA-BB",
        "Tabel 2.2  Pemetaan kompetensi Prakom (Juknis BPS Nomor 2 Tahun 2021)",
        "Tabel 2.3  Jadwal pelaksanaan laboratorium (Gantt Chart)",
        "Tabel 3.1  Kamus data ringkas tabel inti SIBATA-BB",
        "Tabel 3.2  Matriks skenario pengujian black-box",
    ]
    for i, t in enumerate(tables, 1):
        add_text(doc, t, align="left", first=False, after=3, space=1.3)
    doc.add_page_break()

    add_bab(doc, "DAFTAR GAMBAR")
    figs = [
        "Gambar 3.1a  Alur proses bisnis SIBATA-BB",
        "Gambar 3.1b  Diagram use case UML SIBATA-BB",
        "Gambar 3.1c  Diagram kelas UML SIBATA-BB",
        "Gambar 3.1d  Diagram aktivitas UML (swimlane) SIBATA-BB",
        "Gambar 3.2  Entity Relationship Diagram (ERD) kontainer bertingkat",
        "Gambar 3.3  Kamus data dan cuplikan DDL SQL",
        "Gambar 3.4  Mockup antarmuka dashboard web dan menu bot Telegram",
        "Gambar 3.5  Konfigurasi environment server dan webhook SSL Telegram",
        "Gambar 3.6  Wizard bot Telegram: registrasi dan unggah foto",
        "Gambar 3.7  Perintah cepat /pinjam, /kembali, dan /eksekusi",
        "Gambar 3.8  Portal web SIBATA-BB dan generator stiker QR Code",
        "Gambar 3.9  Ekspor buku register ke spreadsheet (.xlsx)",
        "Gambar 3.10  Matriks black-box testing SIBATA-BB",
        "Gambar 3.11  Simulasi pemindaian QR gudang dan sinkronisasi real-time",
        "Gambar 3.12  Sampul dan cuplikan User Manual SIBATA-BB",
    ]
    for t in figs:
        add_text(doc, t, align="left", first=False, after=3, space=1.3)
    doc.add_page_break()


def bab1(doc):
    add_bab(doc, "BAB I")
    add_bab(doc, "PENDAHULUAN")

    add_huruf(doc, "A. Latar Belakang")
    add_text(
        doc,
        "Kejaksaan Negeri Wajo merupakan salah satu instansi penegak hukum di bawah lingkungan Kejaksaan Tinggi Sulawesi Selatan yang berkedudukan di Sengkang, Kabupaten Wajo. Dalam melaksanakan amanat Undang-Undang Republik Indonesia Nomor 11 Tahun 2021 tentang Perubahan atas Undang-Undang Nomor 16 Tahun 2004 tentang Kejaksaan Republik Indonesia, Kejaksaan Negeri Wajo mengemban kekuasaan negara di bidang penuntutan serta kewenangan lain di bidang perdata, tata usaha negara, dan ketertiban umum di wilayah hukum Kabupaten Wajo.",
    )
    add_text(
        doc,
        "Salah satu fungsi vital yang sangat menentukan akuntabilitas peradilan pidana diselenggarakan oleh Seksi Pengelolaan Barang Bukti dan Barang Rampasan (PB3R). Seksi PB3R bertanggung jawab penuh terhadap penerimaan barang bukti penyerahan Tahap II dari penyidik, pengadministrasian buku register perkara, penyimpanan fisik di gudang, pengamanan, pemeliharaan keutuhan fisik, penyediaan barang bukti untuk kepentingan sidang pembuktian di Pengadilan Negeri Sengkang, hingga pelaksanaan eksekusi amar putusan hakim yang telah berkekuatan hukum tetap (inkracht van gewijsde), baik berupa pengembalian kepada pemilik sah, pelelangan untuk kas negara, maupun pemusnahan massal.",
    )
    add_text(
        doc,
        "Dalam kerangka transformasi birokrasi dan modernisasi pelayanan hukum, keberadaan Pejabat Fungsional Pranata Komputer memegang peranan penting dan strategis sebagai motor penggerak digitalisasi instansi (digital enabler). Pranata Komputer bertanggung jawab merencanakan arsitektur TI, menganalisis kebutuhan proses bisnis, merancang basis data relasional, membangun piranti lunak, hingga menguji kehandalan sistem informasi guna memastikan rantai penjagaan barang bukti (chain of custody) terkelola secara transparan, akuntabel, dan terukur.",
    )
    add_text(
        doc,
        "Kondisi tata kelola barang bukti pada Seksi PB3R Kejaksaan Negeri Wajo saat ini telah didukung aplikasi pencatatan di web ruang kantor. Namun, sistem eksisting tersebut masih memiliki kendala mendasar terkait ketergantungan perangkat keras (hardware dependency) dan keterbatasan fleksibilitas mobilitas:",
    )
    add_item(
        doc,
        "1. Ketergantungan pemindai kabel desktop (model kasir). Verifikasi dan pencarian data barang bukti selama ini mengandalkan alat handheld barcode scanner berkabel yang terhubung ke port USB komputer desktop di ruang administrasi PB3R.",
    )
    add_item(
        doc,
        "2. Hambatan luas dan variasi lokasi fisik gudang. Fasilitas penyimpanan terbagi pada beberapa zona fisik terpisah, meliputi brankas barang berharga/uang tunai, lemari loker tersegel untuk narkotika dan barang elektronik, gudang tertutup senjata tajam, serta area lapangan terbuka untuk kendaraan bermotor. Keterbatasan panjang kabel pemindai kasir menyebabkan verifikasi fisik, penataan rak, maupun audit berkala (stock opname) tidak dapat dilakukan langsung di depan fisik barang.",
    )
    add_item(
        doc,
        "3. Inefisiensi alur peminjaman sidang. Ketika Jaksa Penuntut Umum (JPU) memerlukan barang bukti untuk sidang pembuktian di Pengadilan Negeri Sengkang, petugas harus membawa fisik barang ke meja komputer admin atau mencatat nomor perkara secara manual pada buku ekspedisi kertas sebelum data diinput. Hal ini memicu risiko kerusakan fisik serta potensi keterlambatan pemutakhiran buku register.",
    )
    add_item(
        doc,
        "4. Kompleksitas barang bukti multi-item dalam satu wadah segel. Sering dijumpai dalam perkara narkotika atau tindak pidana umum, penyidik melimpahkan barang bukti dalam satu kantong plastik tersegel yang memuat berbagai jenis barang heterogen (misalnya 1 kantong segel berisi 2 paket sabu, 1 timbangan digital, dan 1 unit telepon genggam). Sistem lama mencatat wadah tersebut sebagai satu baris teks tunggal, sehingga petugas kesulitan mencatat mutasi amar putusan yang berbeda pada tiap item (misalnya sabu dimusnahkan, sedangkan telepon genggam dikembalikan kepada saksi).",
    )

    add_huruf(doc, "Penjelasan dan Makna Penamaan Inovasi")
    add_text(
        doc,
        "Untuk menyelesaikan disparitas mobilitas dan kompleksitas data tersebut, dirancang inovasi teknologi bernama SIBATA-BB, akronim dari Sistem Informasi Barang Bukti dan Barang Rampasan. Penamaan SIBATA melambangkan struktur kokoh layaknya tatanan batu bata yang presisi, merefleksikan benteng ketertiban administrasi dan penjagaan barang bukti yang aman, transparan, dan tidak dapat dimanipulasi. Imbuhan -BB menegaskan fokus fungsionalitas sistem pada pengelolaan Barang Bukti dan Barang Rampasan di Kejaksaan Negeri Wajo. SIBATA-BB memadukan mesin bot Telegram berbasis webhook untuk operasional mobile nirkabel di lapangan dengan portal web untuk manajemen pelabelan stiker QR Code kontras tinggi.",
    )
    add_text(
        doc,
        "Guna mengatasi kendala operasional tersebut secara terstruktur dan terukur dalam pelaksanaan Laboratorium Pranata Komputer Tahun 2026, ditetapkan 3 (tiga) rencana kegiatan utama:",
    )
    add_item(doc, "Kegiatan 1: Perancangan Sistem Informasi SIBATA-BB (perancangan arsitektur sistem, basis data relasional model wadah fisik bertingkat, dan antarmuka).")
    add_item(doc, "Kegiatan 2: Pembuatan Sistem Informasi SIBATA-BB (pengembangan Telegram Bot Engine berbasis webhook dan portal web manajemen label QR Code).")
    add_item(doc, "Kegiatan 3: Testing / Pengujian Sistem Informasi SIBATA-BB (uji fungsionalitas black-box, validasi sinkronisasi data seketika Bot–Web, dan simulasi penataan rak gudang).")

    add_huruf(doc, "B. Tujuan")
    add_angka(doc, "1. Tujuan Umum")
    add_text(
        doc,
        "Mewujudkan tata kelola logistik dan administrasi barang bukti yang modern, akuntabel, transparan, dan memiliki mobilitas tinggi di lingkungan Seksi PB3R Kejaksaan Negeri Wajo melalui digitalisasi terpadu berbasis bot perpesanan instan dan portal web (SIBATA-BB).",
    )
    add_angka(doc, "2. Tujuan Khusus")
    add_text(doc, "Tujuan khusus dari masing-masing kegiatan dirinci sebagai berikut:", first=True)
    add_item(
        doc,
        "1. Merancang arsitektur sistem informasi, rancangan antarmuka pengguna (UI/UX), serta skema basis data relasional normalisasi bertingkat (Container & Multi-Item Hierarchy) untuk memisahkan entitas kemasan segel fisik dengan rincian sub-item barang bukti.",
    )
    add_item(
        doc,
        "2. Membangun program aplikasi SIBATA-BB yang terdiri atas Telegram Bot Engine berbasis webhook sebagai alat operasional mobile nirkabel pengganti scanner kabel kasir, serta portal web terpadu untuk pencetakan label stiker QR Code dan modul ekspor data perkara.",
    )
    add_item(
        doc,
        "3. Melaksanakan pengujian sistem informasi secara komprehensif menggunakan metode black-box testing dan simulasi operasional guna memvalidasi ketepatan fungsionalitas fitur, kehandalan transmisi webhook, serta sinkronisasi data dua arah secara real-time.",
    )

    add_huruf(doc, "C. Manfaat")
    add_text(doc, "Pelaksanaan 3 (tiga) kegiatan Laboratorium Pranata Komputer ini memberikan manfaat nyata sebagai berikut.", first=True)
    add_angka(doc, "1. Bagi Satuan Kerja (Seksi PB3R dan Kejaksaan Negeri Wajo)")
    add_item(doc, "a. Menghilangkan ketergantungan pada pemindai kabel kasir desktop tanpa memerlukan anggaran pengadaan perangkat keras scanner barcode nirkabel industri.")
    add_item(doc, "b. Meningkatkan kecepatan dan akurasi petugas dalam mengidentifikasi status perkara serta titik simpan fisik barang bukti hingga ke nomor rak, lemari loker, atau blok parkir melalui pemindaian kamera smartphone.")
    add_item(doc, "c. Menjamin keutuhan rantai penjagaan (chain of custody) melalui pencatatan riwayat sirkulasi pinjam sidang JPU secara real-time, sehingga mencegah selisih data logistik saat persidangan.")
    add_item(doc, "d. Mengakomodasi penanganan eksekusi amar putusan pengadilan yang berbeda pada barang bukti yang berada dalam satu kemasan plastik segel secara tertib dan mandiri.")
    add_angka(doc, "2. Bagi Pejabat Fungsional Pranata Komputer")
    add_item(doc, "a. Mengimplementasikan kompetensi teknis rekayasa piranti lunak, integrasi webhook API, dan perancangan basis data relasional.")
    add_item(doc, "b. Memenuhi kelengkapan bukti fisik butir kegiatan penilaian Angka Kredit Jabatan Fungsional Pranata Komputer Kategori Keahlian (Ahli Pertama) sesuai Peraturan Badan Pusat Statistik Nomor 2 Tahun 2021.")


def bab2(doc):
    add_bab(doc, "BAB II")
    add_bab(doc, "TELAAH MASALAH DAN RENCANA KEGIATAN TI")

    add_huruf(doc, "A. Identifikasi dan Analisis Masalah/Isu TI")
    add_text(
        doc,
        "Analisis isu strategis dilakukan dengan mengidentifikasi kesenjangan (gap) antara kondisi saat ini dan kondisi ideal yang diharapkan pada lingkungan kerja Seksi PB3R Kejaksaan Negeri Wajo. Hasil identifikasi isu dirumuskan ke dalam Tabel 2.1 berikut.",
    )

    table_caption(doc, "Tabel 2.1  Gap analisis isu TI SIBATA-BB")
    rows = [
        ["No", "Aspek isu TI", "Kondisi saat ini", "Kondisi yang diharapkan", "Gap / kesenjangan"],
        [
            "1",
            "Perancangan arsitektur & basis data (Kegiatan 1)",
            "Struktur data barang bukti dalam 1 kantong segel dicatat sebagai satu baris teks tunggal. Pemodelan sistem belum terdokumentasi.",
            "Tersedia rancangan arsitektur dan skema basis data relasional model kontainer bertingkat yang memisahkan wadah fisik dan sub-item barang.",
            "Belum tersedianya dokumen desain sistem, skema ERD bertingkat, kamus data, dan rancangan antarmuka SIBATA-BB.",
        ],
        [
            "2",
            "Ketersediaan sistem & mobilitas input (Kegiatan 2)",
            "Pencatatan dan pemindaian barcode bergantung pada scanner berkabel di meja komputer admin. Peminjaman sidang dicatat manual di buku kertas.",
            "Tersedia aplikasi terpadu: Telegram Bot Engine untuk verifikasi gudang dan pinjam sidang via gawai, serta portal web cetak stiker QR Code.",
            "Ketiadaan sistem terintegrasi klien mobile berbasis webhook dan modul generator stiker QR Code kontras tinggi siap cetak.",
        ],
        [
            "3",
            "Kehandalan & validasi sistem (Kegiatan 3)",
            "Belum pernah dilakukan pengujian formal terhadap kehandalan transmisi webhook nirkabel dan sinkronisasi data real-time.",
            "Sistem teruji fungsionalitasnya (black-box), bebas error kritis, sinkronisasi bot dan web seketika, serta memiliki petunjuk operasional terstandar.",
            "Belum tersusunnya dokumen skenario pengujian (test case), laporan hasil uji, serta buku petunjuk operasional (User Manual).",
        ],
    ]
    t = doc.add_table(rows=len(rows), cols=5)
    fill_table(t, rows, sizes=[1.2, 3.2, 3.8, 3.8, 3.5])

    add_huruf(doc, "B. Deskripsi Kegiatan dan Solusi")
    add_text(
        doc,
        "Guna mengatasi kesenjangan teknis di atas, diusulkan 3 (tiga) rencana kegiatan utama yang saling mendukung.",
    )

    add_angka(doc, "1. Perancangan Sistem Informasi SIBATA-BB (Kegiatan 1)")
    add_text(
        doc,
        "Kegiatan ini mencakup perancangan arsitektur sistem secara menyeluruh, perancangan model proses bisnis, serta perancangan skema basis data relasional model kontainer bertingkat (Container and Multi-Item Hierarchy). Pada tahap ini disusun Entity Relationship Diagram (ERD), Data Flow Diagram (DFD) / Unified Modeling Language (UML), spesifikasi kamus data (Data Dictionary), skrip Data Definition Language (DDL) SQL, serta perancangan antarmuka pengguna (wireframe/mockup) untuk bot Telegram dan dashboard web. Skema ini dirancang khusus untuk memecahkan masalah pencatatan barang bukti multi-item dalam satu kantong segel plastik.",
    )

    add_angka(doc, "2. Pembuatan Sistem Informasi SIBATA-BB (Kegiatan 2)")
    add_text(doc, "Kegiatan ini merupakan implementasi teknis penulisan kode program untuk membangun ekosistem SIBATA-BB yang terdiri dari dua komponen utama.", first=True)
    add_item(
        doc,
        "Mesin Bot Telegram (Webhook Engine). Membangun logika bot interaktif berbasis webhook SSL dengan arsitektur percakapan wizard terarah untuk menangani registrasi Tahap II langsung dari smartphone, pengambilan dan pengunggahan foto fisik barang bukti menggunakan kamera ponsel, perintah cepat mutasi pinjam sidang (/pinjam dan /kembali), eksekusi amar putusan (/eksekusi), serta pengaturan whitelist otorisasi hak akses petugas.",
    )
    add_item(
        doc,
        "Portal Web SIBATA-BB. Membangun aplikasi web administratif menggunakan kerangka kerja Laravel 11 untuk entri data perkara berskala besar, pengelolaan antrean cetak (print queue) stiker label thermal QR Code kontras tinggi, halaman penampil publik (public card view) yang ringan diakses masyarakat/petugas, serta modul ekspor rekapitulasi buku register ke spreadsheet (.xlsx).",
    )

    add_angka(doc, "3. Testing / Pengujian Sistem Informasi SIBATA-BB (Kegiatan 3)")
    add_text(
        doc,
        "Kegiatan ini mencakup pelaksanaan pengujian fungsionalitas secara menyeluruh terhadap seluruh modul piranti lunak SIBATA-BB menggunakan metode black-box testing. Pengujian difokuskan pada respon webhook bot Telegram, akurasi pembacaan kamera ponsel terhadap label QR Code di berbagai kondisi pencahayaan gudang, pengujian sinkronisasi data seketika (two-way real-time data sync) antara ponsel dan dashboard web, serta simulasi alur pinjam-kembali barang bukti persidangan bersama staf PB3R. Seluruh temuan dituangkan ke dalam Laporan Pengujian Sistem Informasi serta dilengkapi Buku Petunjuk Pengoperasian Program Aplikasi (User Manual).",
    )

    add_huruf(doc, "C. Pemetaan Kompetensi Prakom")
    add_text(
        doc,
        "Kegiatan-kegiatan yang diusulkan dipetakan dengan kompetensi Pranata Komputer merujuk pada butir kegiatan dalam Peraturan Badan Pusat Statistik Nomor 2 Tahun 2021 tentang Petunjuk Teknis Penilaian Angka Kredit Jabatan Fungsional Pranata Komputer untuk kategori Keahlian (Jenjang Ahli Pertama).",
    )
    table_caption(doc, "Tabel 2.2  Pemetaan kompetensi Prakom (Juknis BPS Nomor 2 Tahun 2021)")
    rows = [
        ["No", "Rencana kegiatan TI", "Butir kegiatan (Juknis 2021)", "Unsur", "Sub-unsur", "Output / bukti fisik"],
        [
            "1",
            "Perancangan Sistem Informasi SIBATA-BB (Kegiatan 1)",
            "III.A7 Merancang sistem informasi / merancang basis data sistem informasi",
            "Sistem Informasi dan Multimedia",
            "Sistem Informasi",
            "Dokumen perancangan (ERD, kamus data, DDL, mockup UI)",
        ],
        [
            "2",
            "Pembuatan Sistem Informasi SIBATA-BB (Kegiatan 2)",
            "Membuat program aplikasi sistem informasi",
            "Sistem Informasi dan Multimedia",
            "Sistem Informasi",
            "Program aplikasi (source code, screenshot web & bot)",
        ],
        [
            "3",
            "Testing / Pengujian Sistem Informasi SIBATA-BB (Kegiatan 3)",
            "Melakukan pengujian sistem informasi",
            "Sistem Informasi dan Multimedia",
            "Sistem Informasi",
            "Laporan pengujian sistem informasi dan User Manual",
        ],
    ]
    t = doc.add_table(rows=len(rows), cols=6)
    fill_table(t, rows, sizes=[1.0, 3.2, 3.4, 2.4, 2.2, 3.3])

    add_huruf(doc, "D. Penjadwalan (Gantt Chart)")
    add_text(
        doc,
        "Rangkaian kegiatan Laboratorium Pranata Komputer SIBATA-BB dilaksanakan selama 3 (tiga) minggu terhitung mulai tanggal 8 September s.d. 25 September 2026 (mencakup 14 hari kerja efektif). Rincian harian disajikan pada Tabel 2.3. Sel = Selasa, Rab = Rabu, Kam = Kamis, Jum = Jumat, Sen = Senin.",
    )

    table_caption(doc, "Tabel 2.3  Jadwal pelaksanaan laboratorium (Gantt Chart) September 2026")
    days = ["8\nSel", "9\nRab", "10\nKam", "11\nJum", "14\nSen", "15\nSel", "16\nRab", "17\nKam", "18\nJum", "21\nSen", "22\nSel", "23\nRab", "24\nKam", "25\nJum"]
    # marks: index 0..13 for those 14 days
    gantt = [
        ("1. Persiapan, koordinasi, dan analisis kebutuhan", [], ""),
        ("   a. Koordinasi satker & identifikasi masalah scanner kabel", [0], "Notulensi koordinasi"),
        ("   b. Coaching & pengarahan teknis mentor/coach", [0], "Lembar bimbingan"),
        ("2. Perancangan Sistem Informasi SIBATA-BB (Kegiatan 1)", [], ""),
        ("   a. Alur proses bisnis & DFD/UML", [1], "Diagram alur & UML"),
        ("   b. Skema ERD Container & Multi-Item Hierarchy", [2], "Diagram ERD"),
        ("   c. Kamus data & skrip DDL SQL", [3], "Kamus data & DDL"),
        ("   d. Mockup UI dashboard web & state bot Telegram", [3], "Dokumen rancangan"),
        ("3. Pembuatan Sistem Informasi SIBATA-BB (Kegiatan 2)", [], ""),
        ("   a. Environment server, webhook SSL, & koneksi basis data", [4], "Konfigurasi server"),
        ("   b. Coding wizard bot (registrasi, foto, whitelist)", [5, 6], "Source code bot"),
        ("   c. Coding perintah /pinjam, /kembali, /eksekusi", [7, 8], "Source code mutasi"),
        ("   d. Portal web: antrean cetak stiker QR & resolver", [8], "Source code web"),
        ("   e. Integrasi ekspor register .xlsx", [9], "Modul ekspor"),
        ("4. Testing / Pengujian Sistem Informasi (Kegiatan 3)", [], ""),
        ("   a. Skenario uji & black-box testing", [10], "Formulir test case"),
        ("   b. Uji sinkronisasi Bot–Web & simulasi gudang", [11, 12], "Laporan pengujian"),
        ("   c. Penyusunan User Manual", [13], "Buku User Manual"),
        ("5. Penyusunan Laporan Labkom", list(range(14)), "Draf Bab I–IV"),
        ("   a. Penyusunan berkala & evaluasi progres", list(range(14)), "Laporan Labkom"),
    ]
    header = ["Tahapan kegiatan"] + days + ["Output"]
    rows = [header]
    for title, marks, output in gantt:
        row = [title]
        for i in range(14):
            row.append("■" if i in marks else "")
        row.append(output)
        rows.append(row)

    t = doc.add_table(rows=len(rows), cols=16)
    set_table_full(t)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in (0, 15) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(val))
            head = i == 0
            group = row[0].startswith(("1.", "2.", "3.", "4.", "5.")) if i > 0 else False
            set_run(r, size=7 if j not in (0, 15) else 8, bold=head or group, color=WHITE if head else INK)
            if head:
                shade(cell, "0B1F3A")
            elif group:
                shade(cell, "E8EEF6")
            elif j not in (0, 15) and val == "■":
                shade(cell, "C9A227")
            set_cell_border(cell, sz="4")

    add_text(
        doc,
        "Keterangan: sel berwarna emas menandai hari kerja dilaksanakannya tahap tersebut. Kegiatan 5 (penyusunan laporan) berlangsung paralel sepanjang periode laboratorium.",
        first=False,
        before=6,
        size=11,
        italic=True,
    )


def bab3(doc):
    add_bab(doc, "BAB III")
    add_bab(doc, "HASIL KEGIATAN DAN BUKTI KEGIATAN")
    add_huruf(doc, "A. Pelaksanaan dan Hasil Kegiatan Prakom")

    add_angka(doc, "1. Perancangan Sistem Informasi SIBATA-BB (Kegiatan 1)")
    add_text(
        doc,
        "Kegiatan ini dilaksanakan sebagai tahap awal pengembangan sistem informasi SIBATA-BB di Seksi Pengelolaan Barang Bukti dan Barang Rampasan (PB3R) Kejaksaan Negeri Wajo. Perancangan dilakukan untuk menghasilkan cetak biru arsitektur sistem yang mampu mengatasi ketergantungan perangkat pemindai kabel desktop serta mengakomodasi pengelolaan barang bukti multi-item dalam satu wadah kontainer segel. Pelaksanaan kegiatan dilakukan pada tanggal 9–11 September 2026 melalui empat tahapan berikut.",
    )

    add_huruf(doc, "a. Perancangan Alur Proses Bisnis dan Pemodelan Sistem (9 September 2026)")
    add_text(
        doc,
        "Tahap pertama diawali dengan menganalisis proses bisnis eksisting pengelolaan barang bukti, mulai dari penerimaan Tahap II dari penyidik, penyimpanan di gudang, peminjaman untuk sidang di Pengadilan Negeri Sengkang, hingga eksekusi akhir putusan (inkracht). Berdasarkan analisis tersebut, dirancang alur proses bisnis baru (business process reengineering) yang mengintegrasikan bot Telegram dengan portal web. Pemodelan digambarkan menggunakan Unified Modeling Language (UML): use case, class diagram, dan activity diagram ber-swimlane.",
    )
    add_text(doc, "Output: dokumen perancangan alur proses bisnis dan diagram UML sistem SIBATA-BB.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-1a-alur-proses-bisnis.png", "Gambar 3.1a  Diagram alur proses bisnis SIBATA-BB")
    add_figure(doc, "gambar-3-1b-uml-use-case.png", "Gambar 3.1b  Diagram use case UML SIBATA-BB")
    add_figure(doc, "gambar-3-1c-uml-class.png", "Gambar 3.1c  Diagram kelas UML SIBATA-BB")
    add_figure(doc, "gambar-3-1d-uml-activity.png", "Gambar 3.1d  Diagram aktivitas UML (swimlane) SIBATA-BB")

    add_huruf(doc, "b. Perancangan Skema Basis Data Relasional Kontainer dan Multi-Item (10 September 2026)")
    add_text(
        doc,
        "Tahap kedua merancang struktur basis data relasional guna menyelesaikan masalah pencatatan barang bukti campuran/heterogen (misalnya 1 kantong segel plastik yang memuat beberapa jenis barang berbeda). Perancangan menerapkan skema normalisasi bertingkat yang memisahkan entitas kemasan fisik (physical_units) dengan rincian sub-item barang bukti (sip_evidence_items), sehingga setiap sub-item dapat dimutakhirkan status hukum dan eksekusinya secara mandiri tanpa merusak catatan wadah utamanya.",
    )
    add_text(
        doc,
        "Aturan wadah: unit SINGLE memakai kode BB-YYYY-XXX dan stiker menempel langsung pada barang; unit PACK memakai kode PKT-YYYY-XXX dan satu stiker menempel di luar wadah. Isi paket tidak diberi stiker terpisah. Mutasi sirkulasi (pinjam/kembali) mengikuti unit fisik, sedangkan amar putusan mengikuti tiap item.",
    )
    add_text(doc, "Output: rancangan Entity Relationship Diagram (ERD) bertingkat yang memetakan relasi antara tabel perkara, kontainer fisik, sub-item barang bukti, dan riwayat mutasi.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-2-erd.png", "Gambar 3.2  Entity Relationship Diagram (ERD) kontainer bertingkat SIBATA-BB")

    add_huruf(doc, "c. Penyusunan Kamus Data dan Skrip Definisi Basis Data (11 September 2026)")
    add_text(
        doc,
        "Tahap ketiga mencakup penyusunan spesifikasi kamus data (Data Dictionary) yang merinci tipe data, panjang karakter, kunci primer, dan kunci asing untuk setiap tabel. Berdasarkan kamus data tersebut, disusun skrip Data Definition Language (DDL) dalam format SQL untuk menginisialisasi tabel relasional pada basis data SIBATA-BB.",
    )
    table_caption(doc, "Tabel 3.1  Kamus data ringkas tabel inti SIBATA-BB")
    rows = [
        ["Tabel", "Kunci", "Atribut penting", "Fungsi"],
        ["cases", "id / case_number", "defendant_name, prosecutor_name, case_status", "Register perkara Tahap II"],
        ["physical_units", "id / unit_code", "unit_type, storage_location, current_status, is_printed", "Wadah fisik + stiker QR"],
        ["sip_evidence_items", "id / physical_unit_id", "item_name, category, verdict_status, data eksekusi", "Isi barang per unit"],
        ["mutations", "id / physical_unit_id", "mutation_type, borrower_name, handled_by", "Jejak chain of custody"],
        ["telegram_whitelist", "telegram_chat_id", "user_name, role, is_active", "Otorisasi bot lapangan"],
        ["users", "email", "role admin|petugas, is_active", "Login portal web"],
    ]
    t = doc.add_table(rows=len(rows), cols=4)
    fill_table(t, rows, sizes=[3.4, 3.6, 5.0, 3.5])
    add_text(doc, "Output: dokumen kamus data sistem dan berkas skrip DDL SQL siap eksekusi pada database server.", first=True, before=8)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-3-kamus-data-ddl.png", "Gambar 3.3  Struktur kamus data dan cuplikan DDL SQL SIBATA-BB")

    add_huruf(doc, "d. Perancangan Antarmuka Pengguna Dashboard Web dan Bot Telegram (11 September 2026)")
    add_text(
        doc,
        "Tahap keempat merancang antarmuka pengguna (wireframe/mockup) untuk portal web administratif dan struktur menu interaktif (wizard) pada bot Telegram. Perancangan difokuskan pada kemudahan navigasi, tata letak tombol aksi cepat (inline keyboard), serta kejelasan informasi status barang bukti guna memastikan kenyamanan pengguna di lapangan.",
    )
    add_text(doc, "Output: dokumen rancangan antarmuka halaman dashboard web SIBATA-BB dan struktur navigasi bot Telegram.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-4-mockup-ui.png", "Gambar 3.4  Rancangan mockup antarmuka dashboard web dan alur menu bot Telegram SIBATA-BB")

    add_angka(doc, "2. Pembuatan Sistem Informasi SIBATA-BB (Kegiatan 2)")
    add_text(
        doc,
        "Kegiatan ini merupakan implementasi teknis pembangunan piranti lunak SIBATA-BB berdasarkan cetak biru perancangan pada Kegiatan 1. Pembangunan mencakup penyiapan lingkungan server, pengembangan mesin bot Telegram berbasis webhook, pembangunan portal web manajemen label QR Code, hingga modul ekspor data register. Pelaksanaan kegiatan dilakukan pada tanggal 14–21 September 2026 melalui lima tahapan berikut.",
    )

    add_huruf(doc, "a. Penyiapan Infrastruktur Server, Konfigurasi Webhook SSL, dan Koneksi Basis Data (14 September 2026)")
    add_text(
        doc,
        "Tahap awal pembuatan sistem dilakukan dengan menyiapkan lingkungan peladen berbasis kerangka kerja PHP Laravel 11. Dilakukan pengaturan koneksi basis data relasional, konfigurasi keamanan lingkungan (environment variables), serta pengaktifan protokol HTTPS dan pendaftaran URL webhook SSL pada Telegram BotFather agar bot dapat menerima dan mengirimkan pesan secara instan dan aman. Perintah operasional yang digunakan adalah php artisan telegram:set-webhook.",
    )
    add_text(doc, "Output: berkas konfigurasi peladen, pengaturan koneksi basis data, dan status aktif webhook SSL Telegram SIBATA-BB.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-5-webhook-ssl.png", "Gambar 3.5  Konfigurasi environment server dan registrasi webhook SSL Telegram SIBATA-BB")

    add_huruf(doc, "b. Pengkodean Modul Wizard Bot Telegram untuk Registrasi dan Verifikasi Lapangan (15–16 September 2026)")
    add_text(
        doc,
        "Tahap kedua menulis kode program modul interaktif bot Telegram menggunakan arsitektur percakapan wizard. Modul ini memfasilitasi petugas melakukan registrasi data Tahap II, memilih jenis unit SINGLE atau PACK, mengunggah foto fisik barang bukti secara langsung, serta memverifikasi whitelist hak akses petugas di area gudang penyimpanan.",
    )
    add_text(doc, "Output: berkas kode program pengendali bot Telegram untuk registrasi awal dan manajemen verifikasi gudang.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-6-wizard-bot.png", "Gambar 3.6  Pengkodean modul wizard bot Telegram untuk registrasi dan unggah foto barang bukti")

    add_huruf(doc, "c. Pengkodean Perintah Cepat Bot Telegram untuk Mutasi dan Eksekusi Perkara (17–18 September 2026)")
    add_text(
        doc,
        "Tahap ketiga mengembangkan perintah cepat (command handlers) pada bot Telegram untuk mengelola sirkulasi peminjaman barang bukti sidang dan eksekusi putusan pengadilan. Perintah /pinjam, /kembali, dan /eksekusi mencatat mutasi secara seketika, lengkap dengan validasi otomatis terhadap ketersediaan fisik barang dan identitas Jaksa Penuntut Umum pemohon. Pada unit PACK, eksekusi dipilih per item sehingga satu wadah dapat menampung nasib hukum yang berbeda.",
    )
    add_text(doc, "Output: berkas kode program pemrosesan perintah mutasi peminjaman sidang dan eksekusi amar putusan pada bot Telegram.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-7-perintah-mutasi.png", "Gambar 3.7  Pengkodean modul perintah cepat /pinjam, /kembali, dan /eksekusi pada bot Telegram")

    add_huruf(doc, "d. Pembangunan Portal Web SIBATA-BB dan Modul Pencetakan Label QR Code (18 September 2026)")
    add_text(
        doc,
        "Tahap keempat membangun portal web administratif SIBATA-BB. Portal ini dilengkapi dashboard manajemen perkara, sistem pengelola antrean cetak untuk label stiker thermal QR Code kontras tinggi (berukuran kompak yang tahan terhadap debu gudang), serta halaman penampil publik yang ringan diakses guna mendukung transparansi informasi. Unit yang sudah tercetak tetap dapat dicetak ulang dari antrean maupun dari daftar inventaris.",
    )
    add_text(doc, "Output: program aplikasi portal web SIBATA-BB beserta modul generator dan pencetakan stiker QR Code massal.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-8-portal-stiker.png", "Gambar 3.8  Tampilan antarmuka portal web SIBATA-BB dan modul generator stiker QR Code")

    add_huruf(doc, "e. Integrasi Modul Ekspor Data Buku Register ke Format Spreadsheet (21 September 2026)")
    add_text(
        doc,
        "Tahap kelima mengintegrasikan fitur ekspor data otomatis ke dalam portal web. Modul ini memungkinkan petugas administrasi PB3R mencetak buku register perkara dan laporan rekapitulasi bulanan secara instan dalam format spreadsheet (.xlsx) tanpa proses rekapitulasi manual di kertas. Setiap baris ekspor mewakili satu item barang bukti, sehingga wadah paket yang berisi beberapa barang muncul sebagai beberapa baris dengan putusan yang dapat berbeda.",
    )
    add_text(doc, "Output: modul fungsional ekspor data laporan rekapitulasi barang bukti berbasis spreadsheet.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-9-ekspor-xlsx.png", "Gambar 3.9  Antarmuka dan hasil ekspor data buku register ke format spreadsheet .xlsx")

    add_angka(doc, "3. Testing / Pengujian Sistem Informasi SIBATA-BB (Kegiatan 3)")
    add_text(
        doc,
        "Kegiatan ini dilaksanakan untuk menguji keandalan, ketepatan fungsionalitas, serta sinkronisasi data pada sistem SIBATA-BB sebelum dioperasikan secara penuh di Seksi PB3R. Pengujian dilakukan melalui metode pengujian fungsional, simulasi lapangan, serta penyusunan panduan operasional pengguna. Pelaksanaan kegiatan dilakukan pada tanggal 22–25 September 2026 melalui tiga tahapan berikut.",
    )

    add_huruf(doc, "a. Penyusunan Skenario Pengujian dan Pelaksanaan Black-Box Testing (22 September 2026)")
    add_text(
        doc,
        "Tahap pertama menyusun skenario pengujian (test case) untuk seluruh fitur SIBATA-BB. Pengujian menggunakan metode black-box testing guna memastikan seluruh komponen—mulai dari respons bot Telegram, akurasi pemindaian kamera gawai terhadap stiker QR Code, hingga operasional tombol pada portal web—berjalan sesuai spesifikasi fungsional tanpa ditemukan galat kritis.",
    )
    table_caption(doc, "Tabel 3.2  Ringkasan hasil pengujian black-box")
    rows = [
        ["No", "Skenario", "Hasil"],
        ["1", "Halaman login menampilkan identitas SIBATA-BB", "LULUS"],
        ["2", "Pengunjung tanpa sesi dialihkan dari dashboard ke /login", "LULUS"],
        ["3", "Admin dapat membuka cases, units, print-labels, reports, whitelist, users", "LULUS"],
        ["4", "Akun petugas ditolak masuk dashboard administratif", "LULUS"],
        ["5", "Registrasi perkara dengan unit SINGLE dan PACK (multi-item) tersimpan", "LULUS"],
        ["6", "Alur pinjam sidang dan kembali gudang mengubah status unit", "LULUS"],
        ["7", "Lembar cetak dan cetak ulang label memuat kode unit", "LULUS"],
        ["8", "Halaman publik /view/{kode} menampilkan identitas unit", "LULUS"],
        ["9", "Webhook menolak Chat ID di luar whitelist tanpa error fatal", "LULUS"],
        ["10", "Update grup Telegram ditangani aman (HTTP 200)", "LULUS"],
    ]
    t = doc.add_table(rows=len(rows), cols=3)
    fill_table(t, rows, sizes=[1.4, 12.0, 2.1], center_last=True)
    add_text(doc, "Output: dokumen lembar skenario pengujian dan laporan hasil uji fungsional black-box.", first=True, before=8)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-10-blackbox.png", "Gambar 3.10  Matriks skenario pengujian black-box dan hasil eksekusi pengujian fitur SIBATA-BB")

    add_huruf(doc, "b. Uji Sinkronisasi Data Seketika dan Simulasi Penataan Rak Gudang (23–24 September 2026)")
    add_text(
        doc,
        "Tahap kedua dilaksanakan dengan simulasi operasional langsung di lingkungan penyimpanan barang bukti Seksi PB3R. Pengujian difokuskan pada sinkronisasi data dua arah secara real-time antara gawai pintar petugas di lapangan (via bot Telegram) dengan basis data pusat dan dashboard web. Dilakukan pula simulasi audit inventaris (stock opname) dan peminjaman sidang guna memastikan eliminasi ketergantungan pemindai kabel kasir.",
    )
    add_text(doc, "Output: berita acara simulasi lapangan dan dokumentasi keberhasilan sinkronisasi data real-time antara bot Telegram dan portal web.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-11-sinkronisasi.png", "Gambar 3.11  Simulasi pemindaian QR Code via smartphone di gudang dan pengujian sinkronisasi real-time pada dashboard web")

    add_huruf(doc, "c. Penyusunan Buku Petunjuk Pengoperasian Program Aplikasi (25 September 2026)")
    add_text(
        doc,
        "Tahap ketiga menyusun Buku Petunjuk Pengoperasian (User Manual) SIBATA-BB. Panduan ini berisi penjelasan ringkas namun komprehensif mengenai cara registrasi barang bukti melalui bot Telegram, tata cara pencetakan label stiker QR Code, alur peminjaman sidang oleh JPU, hingga pengoperasian dashboard administrasi web bagi petugas PB3R.",
    )
    add_text(doc, "Output: dokumen Buku Petunjuk Pengoperasian (User Manual) sistem SIBATA-BB.", first=True)
    add_text(doc, "Bukti:", first=False, after=2)
    add_figure(doc, "gambar-3-12-user-manual.png", "Gambar 3.12  Sampul dan cuplikan isi dokumen Buku Petunjuk Pengoperasian (User Manual) SIBATA-BB")

    add_huruf(doc, "B. Dampak Hasil Kegiatan")
    add_angka(doc, "1. Perancangan Sistem Informasi SIBATA-BB (Kegiatan 1)")
    add_text(
        doc,
        "Dampak positif: tersedianya cetak biru arsitektur sistem dan model basis data relasional bertingkat (Container & Multi-Item Hierarchy) yang memberikan kejelasan struktural dalam mengelola barang bukti multi-item dalam satu kantong segel. Hal ini menghilangkan ambiguitas pencatatan data dan menjadi fondasi kokoh bagi pembangunan aplikasi yang terstandar.",
    )
    add_angka(doc, "2. Pembuatan Sistem Informasi SIBATA-BB (Kegiatan 2)")
    add_text(
        doc,
        "Dampak positif: terwujudnya program aplikasi SIBATA-BB secara utuh yang mengintegrasikan bot Telegram berbasis webhook dengan portal web manajemen QR Code. Inovasi ini berhasil menghapus ketergantungan pada pemindai kabel desktop (hardware dependency), memberikan mobilitas penuh bagi petugas untuk memverifikasi fisik barang langsung di depan rak penyimpanan gudang, serta mempercepat proses rekapitulasi buku register secara digital.",
    )
    add_angka(doc, "3. Testing / Pengujian Sistem Informasi SIBATA-BB (Kegiatan 3)")
    add_text(
        doc,
        "Dampak positif: memastikan keandalan operasional sistem melalui pengujian black-box dan simulasi lapangan yang membuktikan bahwa transmisi data bot dan web berjalan akurat, cepat, dan bebas galat. Adanya Buku Petunjuk Pengoperasian (User Manual) juga menjamin keberlanjutan penggunaan sistem oleh pegawai di satuan kerja secara mandiri.",
    )


def bab4(doc):
    add_bab(doc, "BAB IV")
    add_bab(doc, "KENDALA DAN RENCANA TINDAK LANJUT")

    add_huruf(doc, "A. Kendala")
    add_angka(doc, "1. Perancangan Sistem Informasi SIBATA-BB (Kegiatan 1)")
    add_item(doc, "• Pemodelan hubungan relasional antara kemasan kontainer fisik dan sub-item barang bukti yang heterogen memerlukan ketelitian ekstra agar aturan normalisasi basis data tetap terpenuhi secara optimal.")
    add_item(doc, "• Penyesuaian batasan alur proses bisnis harus diselaraskan secara saksama dengan Standar Operasional Prosedur (SOP) pengelolaan barang bukti yang berlaku di lingkungan Kejaksaan Republik Indonesia.")

    add_angka(doc, "2. Pembuatan Sistem Informasi SIBATA-BB (Kegiatan 2)")
    add_item(doc, "• Pengaturan konfigurasi server web untuk mendukung komunikasi webhook Telegram memerlukan penyesuaian sertifikat SSL yang valid agar transmisi data interaktif berjalan tanpa hambatan keamanan.")
    add_item(doc, "• Penyesuaian tata letak cetak stiker QR Code memerlukan uji coba berulang agar ukuran stiker pas saat dicetak pada kertas label thermal dan tetap terbaca jelas meskipun ditempel di permukaan yang terkena debu gudang.")

    add_angka(doc, "3. Testing / Pengujian Sistem Informasi SIBATA-BB (Kegiatan 3)")
    add_item(doc, "• Kondisi pencahayaan yang minim di beberapa sudut lorong gudang penyimpanan barang bukti sempat memengaruhi kecepatan fokus kamera gawai pintar dalam memindai label QR Code pada tahap uji coba awal.")
    add_item(doc, "• Keterbatasan waktu pelaksanaan laboratorium komputer menuntut pengelolaan jadwal pengujian dan simulasi lapangan yang sangat disiplin bersama staf Seksi PB3R.")

    add_huruf(doc, "B. Rencana Tindak Lanjut")
    add_angka(doc, "1. Perancangan Sistem Informasi SIBATA-BB (Kegiatan 1)")
    add_text(
        doc,
        "Melakukan evaluasi berkala terhadap struktur basis data seiring dengan dinamika penambahan jenis perkara atau karakteristik barang bukti khusus di masa mendatang.",
    )
    add_angka(doc, "2. Pembuatan Sistem Informasi SIBATA-BB (Kegiatan 2)")
    add_text(
        doc,
        "Mengembangkan fitur tambahan pada portal web berupa notifikasi otomatis (automatic reminder) berbasis pesan instan untuk mengingatkan jadwal pengembalian barang bukti pinjaman sidang yang melewati batas waktu.",
    )
    add_angka(doc, "3. Testing / Pengujian Sistem Informasi SIBATA-BB (Kegiatan 3)")
    add_text(
        doc,
        "Menyelenggarakan sesi sosialisasi dan bimbingan teknis (coaching clinic) singkat bagi seluruh Jaksa Penuntut Umum dan staf pengelola PB3R Kejaksaan Negeri Wajo agar pemanfaatan SIBATA-BB dapat berjalan secara optimal dan berkesinambungan.",
    )

    add_text(doc, "Demikian laporan laboratorium ini disusun sebagai pertanggungjawaban pelaksanaan tiga kegiatan Pranata Komputer serta sebagai bukti fisik penilaian angka kredit jabatan fungsional.", first=True, before=18)


def main():
    missing = [p.name for p in [
        GAMBAR / "gambar-3-1a-alur-proses-bisnis.png",
        GAMBAR / "gambar-3-12-user-manual.png",
    ] if not p.exists()]
    if missing:
        raise SystemExit(f"Gambar belum lengkap: {missing}. Jalankan _bangun_gambar.py terlebih dahulu.")

    doc = setup_doc()
    cover(doc)
    toc(doc)
    bab1(doc)
    bab2(doc)
    bab3(doc)
    bab4(doc)
    doc.save(OUT)
    print("saved", OUT, "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    main()
